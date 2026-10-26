"""Hàm tiện ích tạo báo cáo .docx bằng python-docx (Times New Roman 13, giống các báo cáo lab trước)."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"


class Report:
    def __init__(self):
        self.doc = Document()
        self.n_fig = 0
        self.n_tab = 0
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)  # A4
        s.left_margin, s.right_margin = Cm(3.0), Cm(2.0)
        s.top_margin, s.bottom_margin = Cm(2.0), Cm(2.0)
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(13)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.2
        for lvl, size in ((1, 16), (2, 14), (3, 13)):
            h = self.doc.styles[f"Heading {lvl}"]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
            h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                h.element.rPr.rFonts.attrib.pop(qn(attr), None)  # bỏ font theme để dùng Times New Roman
            h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
            h.paragraph_format.space_after = Pt(4)
        self._page_numbers()

    # ------------------------------------------------------------------ cơ bản
    def _page_numbers(self):
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(p, "PAGE")

    @staticmethod
    def _field(paragraph, code: str):
        run = paragraph.add_run()
        for tag, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            elif text == code:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = code
            else:
                el = OxmlElement("w:t")
                el.text = text
            run._r.append(el)

    def toc(self):
        p = self.doc.add_paragraph()
        self._field(p, 'TOC \\o "1-3" \\h \\z \\u')  # số trang được cập nhật bởi report/finalize_word.ps1

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text):
        return self.doc.add_heading(text, level=1)

    def h2(self, text):
        return self.doc.add_heading(text, level=2)

    def h3(self, text):
        return self.doc.add_heading(text, level=3)

    def para(self, text: str = "", bold=False, italic=False, size=None, align=None, justify=True):
        """Đoạn văn; hỗ trợ **in đậm** inline."""
        p = self.doc.add_paragraph()
        parts = text.split("**")
        for i, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run(part)
            r.bold = bold or (i % 2 == 1)
            r.italic = italic
            if size:
                r.font.size = Pt(size)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def bullets(self, items: list[str], style="List Bullet"):
        for it in items:
            p = self.doc.add_paragraph(style=style)
            for i, part in enumerate(it.split("**")):
                if part:
                    p.add_run(part).bold = i % 2 == 1
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def numbered(self, items: list[str]):
        """Danh sách đánh số 1, 2, 3… (bắt đầu lại từ 1 ở mỗi danh sách)."""
        for k, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.add_run(f"{k}.	")
            for i, part in enumerate(it.split("**")):
                if part:
                    p.add_run(part).bold = i % 2 == 1
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def code(self, text: str):
        for line in text.strip("\n").splitlines():
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)

    # ------------------------------------------------------------------ bảng & hình
    def caption(self, kind: str, text: str):
        if kind == "Bảng":
            self.n_tab += 1
            n = self.n_tab
        else:
            self.n_fig += 1
            n = self.n_fig
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{kind} {n}. {text}")
        r.italic = True
        r.font.size = Pt(11)
        return n

    def table(self, df: pd.DataFrame, caption: str | None = None, font=10, col_widths_cm: list[float] | None = None,
              fmt: dict | None = None, decimals: int = 4):
        if caption:
            self.caption("Bảng", caption)
        fmt = fmt or {}
        t = self.doc.add_table(rows=1, cols=len(df.columns))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, c in enumerate(df.columns):
            cell = t.rows[0].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(c))
            r.bold = True
            r.font.size = Pt(font)
            self._shade(cell, "D9E2F3")
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for j, c in enumerate(df.columns):
                v = row[c]
                if c in fmt:
                    s = fmt[c](v)
                elif isinstance(v, float):
                    s = "–" if pd.isna(v) else f"{v:.{decimals}f}"
                else:
                    s = str(v)
                cells[j].text = ""
                r = cells[j].paragraphs[0].add_run(s)
                r.font.size = Pt(font)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if col_widths_cm:
            for row in t.rows:
                for j, w in enumerate(col_widths_cm):
                    row.cells[j].width = Cm(w)
        self.doc.add_paragraph()
        return t

    @staticmethod
    def _shade(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def figure(self, path: Path, caption: str, width_cm: float = 15.5):
        if not Path(path).exists():
            self.para(f"[Thiếu hình: {Path(path).name}]", italic=True)
            return
        self.doc.add_picture(str(path), width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.caption("Hình", caption)

    def save(self, path: Path):
        self.doc.save(str(path))
